# -*- coding: utf-8 -*-
"""
BỘ KIỂM THỬ TỰ ĐỘNG TOÀN DIỆN — TÍNH NĂNG NHẬP DANH SÁCH TÀI KHOẢN HÀNG LOẠT (EXCEL/CSV)
ĐÁP ỨNG ĐẦY ĐỦ 6 TIÊU CHÍ NGHIỆM THU:
1. Nhập file mẫu 10 dòng (có 2 dòng lỗi cố ý) -> 8 thành công, 2 báo lỗi đúng lý do, các dòng đúng không bị ảnh hưởng.
2. Mật khẩu mỗi người khác nhau 100%, đăng nhập được ngay sau khi nhập.
3. File kết quả tải về có đủ họ tên + tên đăng nhập + mật khẩu.
4. Quản trị trường A không nhập được cho trường B (RBAC bảo vệ dữ liệu trường).
5. Chạy lại cùng file -> Chống trùng lặp (Idempotency), không tạo tài khoản trùng.
6. Hỗ trợ cả file Excel (.xlsx) và file CSV (.csv).
"""

import os
import io
import sys
import time
import secrets
import unittest
from pathlib import Path
from werkzeug.security import generate_password_hash, check_password_hash
import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

from app import app, get_db, init_db, DEMO_SCHOOL_ID


class TestPromptImportUsers(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        app.config["WTF_CSRF_ENABLED"] = False
        with app.app_context():
            init_db()

            # Tạo tài khoản admin test cho trường 1 và trường 2
            db = get_db()
            cur = db.cursor()
            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'admin_truong_1'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, vai_tro, truong_id, trang_thai, so_du_gio, email)
                    VALUES ('admin_truong_1', 'Quản trị viên Trường 1', ?, 'school_admin', 1, 'hoat_dong', 100.0, 'admin1@uka.edu.vn')
                """, (generate_password_hash("Admin1Pass123!"),))

            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'admin_truong_2'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, vai_tro, truong_id, trang_thai, so_du_gio, email)
                    VALUES ('admin_truong_2', 'Quản trị viên Trường 2', ?, 'school_admin', 2, 'hoat_dong', 100.0, 'admin2@thuoc.edu.vn')
                """, (generate_password_hash("Admin2Pass123!"),))
            db.commit()

    def setUp(self):
        self.app_context = app.app_context()
        self.app_context.push()
        self.client = app.test_client()

    def tearDown(self):
        self.app_context.pop()

    def login_as(self, username, password):
        self.client.get("/logout")
        res = self.client.post("/login", data={
            "ma_hoc_sinh": username,
            "mat_khau": password
        }, follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        return res

    def make_unique_phone(self):
        return "09" + "".join(secrets.choice("0123456789") for _ in range(8))

    def make_unique_email(self, prefix="user"):
        return f"{prefix}.{secrets.token_hex(6)}@uka.test"

    # =========================================================================
    # TEST 1: TẢI FILE EXCEL MẪU
    # =========================================================================
    def test_01_download_template(self):
        """[TC 1]: Tải file Excel mẫu (/admin/import-users/template) trả về đúng cấu trúc 6 cột."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        res = self.client.get("/admin/import-users/template")
        self.assertEqual(res.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res.content_type)

        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Mau_Nhap_Tai_Khoan", wb.sheetnames)
        ws = wb["Mau_Nhap_Tai_Khoan"]

        headers = [cell.value for cell in ws[1]]
        expected_headers = ["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"]
        self.assertEqual(headers, expected_headers, "Header file mẫu phải đúng 6 cột quy định")
        self.assertGreaterEqual(ws.max_row, 4, "File mẫu phải có ít nhất 3 dòng dữ liệu minh họa")
        print("\n[PASS - TC 1]: Tải file Excel mẫu thành công, chuẩn 6 cột quy định.")

    # =========================================================================
    # TEST 2: NHẬP 10 DÒNG (2 LỖI CỐ Ý) -> 8 THÀNH CÔNG, 2 BÁO LỖI (TIÊU CHÍ 1)
    # =========================================================================
    def test_02_import_10_rows_with_2_errors(self):
        """[TIÊU CHÍ 1]: Nhập file mẫu 10 dòng (2 lỗi cố ý) -> 8 thành công, 2 báo lỗi đúng lý do."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh_Sach"
        ws.append(["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"])

        email1 = self.make_unique_email("hs.an")
        email2 = self.make_unique_email("hs.binh")
        email3 = self.make_unique_email("gv.cuong")

        test_rows = [
            ["Nguyễn Minh An", "hocsinh", "10A1", email1, self.make_unique_phone(), "Học sinh 1"],
            ["Trần Bảo Bình", "hocsinh", "10A1", email2, self.make_unique_phone(), "Học sinh 2"],
            ["Lê Quốc Cường", "giaovien", "Tổ Toán", email3, self.make_unique_phone(), "Giáo viên Toán"],
            ["Phạm Thùy Dung", "quantruong", "Ban Giám Hiệu", self.make_unique_email("qt.dung"), self.make_unique_phone(), "Phó Hiệu Trưởng"],
            ["Hoàng Tuấn Em", "hocsinh", "10A2", self.make_unique_email("hs.em"), self.make_unique_phone(), "Học sinh 3"],
            ["", "hocsinh", "10A2", self.make_unique_email("err1"), self.make_unique_phone(), "Dòng lỗi 1: Họ tên trống"],
            ["Đặng Hồng Phúc", "vai_tro_sai_quy_dinh", "10A2", self.make_unique_email("err2"), self.make_unique_phone(), "Dòng lỗi 2: Vai trò sai"],
            ["Vũ Thị Giang", "hocsinh", "10A3", self.make_unique_email("hs.giang"), self.make_unique_phone(), "Học sinh 4"],
            ["Bùi Trọng Hải", "giaovien", "Tổ Lý", self.make_unique_email("gv.hai"), self.make_unique_phone(), "Giáo viên Lý"],
            ["Ngô Mỹ Kim", "hocsinh", "10A3", self.make_unique_email("hs.kim"), self.make_unique_phone(), "Học sinh 5"],
        ]
        for r in test_rows:
            ws.append(r)

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        data = {
            "file": (excel_buf, "danh_sach_10_nguoi.xlsx"),
            "truong_id": "1"
        }
        res = self.client.post("/admin/import-users", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        with self.client.session_transaction() as sess:
            result = sess.get("last_import_result")
            self.assertIsNotNone(result, "Phải có kết quả import lưu trong session")
            self.assertEqual(result["success_count"], 8, "Phải có đúng 8 dòng thành công")
            self.assertEqual(result["error_count"], 2, "Phải có đúng 2 dòng bị báo lỗi")
            self.assertEqual(result["total_rows"], 10, "Tổng số dòng xử lý phải là 10")

            errors = result["errors"]
            self.assertEqual(len(errors), 2)
            self.assertIn("Họ và tên không được để trống", errors[0]["ly_do"])
            self.assertIn("không hợp lệ", errors[1]["ly_do"])

        # Kiểm tra CSDL: 8 người phải được tạo và ở trạng thái hoat_dong
        db = get_db()
        cur = db.cursor()
        cur.execute("SELECT * FROM users WHERE email IN (?, ?, ?)", (email1, email2, email3))
        inserted_users = cur.fetchall()
        self.assertEqual(len(inserted_users), 3)
        for u in inserted_users:
            self.assertEqual(u["trang_thai"], "hoat_dong")
            self.assertEqual(u["truong_id"], 1)
            if u["vai_tro"] == "hoc_sinh":
                self.assertEqual(u["so_du_gio"], 2.0, "Học sinh phải được cấp đúng 2.0 giờ ban đầu")

        print("\n[PASS - TC 2]: Nhập 10 dòng (2 lỗi cố ý) -> 8 thành công, 2 lỗi chính xác lý do, 8 dòng đúng không bị ảnh hưởng.")

    # =========================================================================
    # TEST 3: MẬT KHẨU MỖI NGƯỜI KHÁC NHAU & ĐĂNG NHẬP NGAY (TIÊU CHÍ 2)
    # =========================================================================
    def test_03_passwords_unique_and_immediate_login(self):
        """[TIÊU CHÍ 2]: Mật khẩu mỗi người khác nhau 100%, đăng nhập được ngay sau khi nhập."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"])
        test_rows = [
            ["Trịnh Minh Khang", "hocsinh", "11A1", self.make_unique_email("khang"), self.make_unique_phone(), "HS 1"],
            ["Đỗ Thanh Long", "hocsinh", "11A1", self.make_unique_email("long"), self.make_unique_phone(), "HS 2"],
            ["Nguyễn Thị Mai", "giaovien", "Tổ Sử", self.make_unique_email("mai"), self.make_unique_phone(), "GV Sử"],
            ["Hoàng Ngọc Nam", "quantruong", "BGH", self.make_unique_email("nam"), self.make_unique_phone(), "Quản trị"],
            ["Phan Thu Oanh", "hocsinh", "11A2", self.make_unique_email("oanh"), self.make_unique_phone(), "HS 3"],
        ]
        for r in test_rows:
            ws.append(r)

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        data = {
            "file": (excel_buf, "test_passwords.xlsx"),
            "truong_id": "1"
        }
        res = self.client.post("/admin/import-users", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        with self.client.session_transaction() as sess:
            result = sess.get("last_import_result")
            self.assertIsNotNone(result)
            success_list = result["success_list"]

        self.assertEqual(len(success_list), 5)

        # 1. Kiểm tra tính ngẫu nhiên độc nhất: Không mật khẩu nào trùng nhau
        passwords = [item["mat_khau"] for item in success_list]
        self.assertEqual(len(set(passwords)), 5, "100% mật khẩu được sinh phải khác nhau riêng biệt!")

        for pwd in passwords:
            self.assertNotEqual(pwd, "uka123")
            self.assertGreaterEqual(len(pwd), 8, "Độ dài mật khẩu tối thiểu 8 ký tự")

        # 2. Kiểm tra đăng nhập được ngay với tài khoản và mật khẩu vừa cấp
        for item in success_list:
            username = item["ma_hoc_sinh"]
            raw_password = item["mat_khau"]

            self.client.get("/logout")
            res_login = self.client.post("/login", data={
                "ma_hoc_sinh": username,
                "mat_khau": raw_password
            }, follow_redirects=True)
            self.assertEqual(res_login.status_code, 200, f"Đăng nhập tài khoản {username} thất bại!")

            with self.client.session_transaction() as user_sess:
                self.assertEqual(user_sess.get("ma_hoc_sinh"), username)
                self.assertEqual(user_sess.get("vai_tro"), item["vai_tro"])

        print("\n[PASS - TC 3]: 100% mật khẩu ngẫu nhiên riêng biệt, đăng nhập được ngay lập tức.")

    # =========================================================================
    # TEST 4: TẢI FILE KẾT QUẢ ĐỦ HỌ TÊN + USERNAME + PASSWORD (TIÊU CHÍ 3)
    # =========================================================================
    def test_04_download_result_excel(self):
        """[TIÊU CHÍ 3]: File kết quả tải về có đủ họ tên, tên đăng nhập, mật khẩu ban đầu."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        wb_in = openpyxl.Workbook()
        ws_in = wb_in.active
        ws_in.append(["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"])
        ws_in.append(["Võ Quốc Phong", "hocsinh", "12A1", self.make_unique_email("phong"), self.make_unique_phone(), "HS 12"])
        ws_in.append(["Trương Mỹ Quyên", "giaovien", "Tổ Hóa", self.make_unique_email("quyen"), self.make_unique_phone(), "GV Hóa"])

        excel_buf = io.BytesIO()
        wb_in.save(excel_buf)
        excel_buf.seek(0)

        self.client.post("/admin/import-users", data={
            "file": (excel_buf, "test_download.xlsx"),
            "truong_id": "1"
        }, content_type="multipart/form-data", follow_redirects=True)

        res = self.client.get("/admin/import-users/download-result")
        self.assertEqual(res.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res.content_type)

        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Danh_Sach_Tai_Khoan", wb.sheetnames)
        ws = wb["Danh_Sach_Tai_Khoan"]

        headers = [cell.value for cell in ws[1]]
        self.assertIn("Họ và tên", headers)
        self.assertIn("Tên đăng nhập (Mã định danh)", headers)
        self.assertIn("Mật khẩu ban đầu", headers)
        self.assertIn("Vai trò", headers)

        data_rows = list(ws.iter_rows(values_only=True))[1:]
        self.assertEqual(len(data_rows), 2, "File kết quả phải chứa đúng 2 dòng tài khoản thành công")

        idx_name = headers.index("Họ và tên")
        idx_user = headers.index("Tên đăng nhập (Mã định danh)")
        idx_pass = headers.index("Mật khẩu ban đầu")

        for r in data_rows:
            self.assertTrue(r[idx_name], "Họ tên không được rỗng")
            self.assertTrue(r[idx_user], "Tên đăng nhập không được rỗng")
            self.assertTrue(r[idx_pass], "Mật khẩu không được rỗng")

        print("\n[PASS - TC 4]: File kết quả tải về chứa đầy đủ họ tên, username, password ban đầu.")

    # =========================================================================
    # TEST 5: QUẢN TRỊ TRƯỜNG A KHÔNG NHẬP ĐƯỢC CHO TRƯỜNG B (TIÊU CHÍ 4)
    # =========================================================================
    def test_05_school_admin_cross_tenant_blocked(self):
        """[TIÊU CHÍ 4]: Quản trị trường 1 gửi request nhập cho trường 2 -> Bị chặn, không tạo cho trường 2."""
        self.login_as("admin_truong_1", "Admin1Pass123!")  # Admin trường 1
        cross_email = self.make_unique_email("cross")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"])
        ws.append(["Hacker Trường", "hocsinh", "10A1", cross_email, self.make_unique_phone(), "Cố ý nhập trường khác"])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        # Cố ý gửi truong_id = 2 (Trường khác)
        data = {
            "file": (excel_buf, "attack_cross_tenant.xlsx"),
            "truong_id": "2"
        }
        res = self.client.post("/admin/import-users", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        html = res.data.decode("utf-8")
        self.assertIn("Quản trị trường chỉ có quyền nhập danh sách cho trường của mình", html)

        # Xác minh trong DB không có tài khoản nào được tạo cho trường 2 từ request này
        db = get_db()
        cur = db.cursor()
        cur.execute("SELECT id FROM users WHERE email = ?", (cross_email,))
        self.assertIsNone(cur.fetchone(), "Tài khoản cross-tenant tuyệt đối không được phép tạo")

        print("\n[PASS - TC 5]: RBAC đa trường hoạt động chuẩn xác: Quản trị trường 1 bị chặn khi nhập cho trường 2.")

    # =========================================================================
    # TEST 6: CHỐNG TRÙNG LẶP (IDEMPOTENCY) — CHẠY LẠI KHÔNG TẠO TRÙNG (TIÊU CHÍ 5)
    # =========================================================================
    def test_06_idempotency_no_duplicate_accounts(self):
        """[TIÊU CHÍ 5]: Chạy lại cùng file cũ -> Báo lỗi trùng lặp, không tạo thêm tài khoản trùng."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        email_a = self.make_unique_email("idem.a")
        email_b = self.make_unique_email("idem.b")
        phone_a = self.make_unique_phone()
        phone_b = self.make_unique_phone()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["họ_tên", "vai_trò", "lớp", "email", "số_điện_thoại", "ghi_chú"])
        ws.append(["Nguyễn Minh An", "hocsinh", "10A1", email_a, phone_a, "Học sinh 1"])
        ws.append(["Trần Bảo Bình", "hocsinh", "10A1", email_b, phone_b, "Học sinh 2"])

        excel_buf1 = io.BytesIO()
        wb.save(excel_buf1)
        excel_buf1.seek(0)

        # Lần 1: Nhập lần đầu thành công
        res1 = self.client.post("/admin/import-users", data={
            "file": (excel_buf1, "chay_lai_file_cu.xlsx"),
            "truong_id": "1"
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res1.status_code, 200)

        with self.client.session_transaction() as sess1:
            self.assertEqual(sess1.get("last_import_result")["success_count"], 2)

        db = get_db()
        cur = db.cursor()
        cur.execute("SELECT COUNT(*) FROM users")
        count_after_first_run = cur.fetchone()[0]

        # Lần 2: Nạp lại file y hệt -> Phải phát hiện trùng lặp 100%
        excel_buf2 = io.BytesIO()
        wb.save(excel_buf2)
        excel_buf2.seek(0)

        res2 = self.client.post("/admin/import-users", data={
            "file": (excel_buf2, "chay_lai_file_cu.xlsx"),
            "truong_id": "1"
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res2.status_code, 200)

        with self.client.session_transaction() as sess2:
            result2 = sess2.get("last_import_result")
            self.assertEqual(result2["success_count"], 0, "Không được tạo thêm tài khoản nào vì đã tồn tại")
            self.assertEqual(result2["error_count"], 2, "Cả 2 dòng đều phải báo lỗi trùng lặp")
            self.assertIn("đã tồn tại trong hệ thống", result2["errors"][0]["ly_do"])

        # Số lượng tài khoản trong DB không đổi
        cur.execute("SELECT COUNT(*) FROM users")
        count_after_second_run = cur.fetchone()[0]
        self.assertEqual(count_after_first_run, count_after_second_run, "Số lượng user trong DB không được tăng khi chạy lại")

        print("\n[PASS - TC 6]: Cơ chế Idempotency chống trùng lặp hoạt động hoàn hảo: chạy lại file cũ không tạo tài khoản trùng.")

    # =========================================================================
    # TEST 7: HỖ TRỢ ĐỊNH DẠNG FILE CSV
    # =========================================================================
    def test_07_import_csv_format(self):
        """[TC 7]: Hỗ trợ nhập danh sách qua file CSV (.csv) định dạng UTF-8."""
        self.login_as("admin_truong_1", "Admin1Pass123!")

        csv_content = (
            "họ_tên,vai_trò,lớp,email,số_điện_thoại,ghi_chú\n"
            f"Đoàn Văn Hậu,hocsinh,11B1,{self.make_unique_email('hau')},{self.make_unique_phone()},Học sinh CSV 1\n"
            f"Trịnh Thùy Linh,giaovien,Tổ Văn,{self.make_unique_email('linh')},{self.make_unique_phone()},Giáo viên Văn CSV\n"
        ).encode("utf-8")

        data = {
            "file": (io.BytesIO(csv_content), "danh_sach_hoc_sinh.csv"),
            "truong_id": "1"
        }
        res = self.client.post("/admin/import-users", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        with self.client.session_transaction() as sess:
            result = sess.get("last_import_result")
            self.assertEqual(result["success_count"], 2)
            self.assertEqual(result["error_count"], 0)

        print("\n[PASS - TC 7]: Hỗ trợ định dạng file CSV chuẩn UTF-8 thành công.")


if __name__ == "__main__":
    unittest.main()
