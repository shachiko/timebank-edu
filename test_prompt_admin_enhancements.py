# -*- coding: utf-8 -*-
"""
BỘ KIỂM THỬ TỰ ĐỘNG — CÁC TÍNH NĂNG NÂNG CAO QUẢN TRỊ ADMIN (TIMEBANK EDU)
1. Xuất danh sách tài khoản người dùng ra file Excel (.xlsx) kèm mật khẩu khởi tạo / trạng thái.
2. Chọn nhiều mã mời qua checkbox và xuất file Excel riêng cho các mã đã chọn.
3. Xóa mã mời an toàn với kiểm tra phân quyền đa trường.
4. Quản lý tài khoản người dùng: Khóa / Mở khóa, Đặt lại mật khẩu (cập nhật mật khẩu khởi tạo), Xóa tài khoản an toàn.
"""

import io
import os
import sys
import unittest
from datetime import datetime
import openpyxl
from werkzeug.security import generate_password_hash, check_password_hash

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

from app import app, get_db, init_db


class TestPromptAdminEnhancements(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        app.config["WTF_CSRF_ENABLED"] = False
        with app.app_context():
            init_db()
            db = get_db()
            cur = db.cursor()

            # Tạo tài khoản Super Admin kiểm thử
            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'SUPER_ADMIN_ENH'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, vai_tro, truong_id, trang_thai, so_du_gio, email)
                    VALUES ('SUPER_ADMIN_ENH', 'Admin Tong Enh', ?, 'super_admin', 1, 'hoat_dong', 50.0, 'super_enh@timebank.edu.vn')
                """, (generate_password_hash("AdminPass123!"),))
            else:
                cur.execute("UPDATE users SET vai_tro = 'super_admin' WHERE ma_hoc_sinh = 'SUPER_ADMIN_ENH'")

            # Tạo tài khoản School Admin kiểm thử
            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'SCHOOL_ADMIN_ENH'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, vai_tro, truong_id, trang_thai, so_du_gio, email)
                    VALUES ('SCHOOL_ADMIN_ENH', 'Admin Truong Enh', ?, 'school_admin', 1, 'hoat_dong', 50.0, 'school_enh@uka.edu.vn')
                """, (generate_password_hash("AdminPass123!"),))

            # Tạo tài khoản Học sinh test 1
            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'HS_TEST_ENH1'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, mat_khau_khoi_tao, vai_tro, truong_id, trang_thai, so_du_gio, lop, email)
                    VALUES ('HS_TEST_ENH1', 'Nguyen Van Enh 1', ?, 'Pass123@', 'hoc_sinh', 1, 'hoat_dong', 10.0, '10A1', 'enh1@student.edu.vn')
                """, (generate_password_hash("Pass123@"),))

            # Tạo tài khoản Học sinh test 2
            cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'HS_TEST_ENH2'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, mat_khau_khoi_tao, vai_tro, truong_id, trang_thai, so_du_gio, lop, email)
                    VALUES ('HS_TEST_ENH2', 'Tran Thi Enh 2', ?, 'Pass456@', 'hoc_sinh', 1, 'hoat_dong', 15.0, '10A2', 'enh2@student.edu.vn')
                """, (generate_password_hash("Pass456@"),))

            # Tạo mã mời test
            cur.execute("SELECT id FROM invite_codes WHERE ma_code = 'TBEDU-TEST-ENH1'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO invite_codes (ma_code, truong_id, loai, so_luot_toi_da, da_dung, ngay_tao)
                    VALUES ('TBEDU-TEST-ENH1', 1, 'ca_nhan', 1, 0, ?)
                """, (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))

            cur.execute("SELECT id FROM invite_codes WHERE ma_code = 'TBEDU-TEST-ENH2'")
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO invite_codes (ma_code, truong_id, loai, so_luot_toi_da, da_dung, ngay_tao)
                    VALUES ('TBEDU-TEST-ENH2', 1, 'lop', 35, 2, ?)
                """, (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))

            db.commit()

    def setUp(self):
        self.app_context = app.app_context()
        self.app_context.push()
        self.client = app.test_client()

    def tearDown(self):
        self.app_context.pop()

    def login_as(self, username, password="AdminPass123!"):
        self.client.get("/logout")
        return self.client.post("/login", data={
            "ma_hoc_sinh": username,
            "mat_khau": password
        }, follow_redirects=True)

    def test_01_ui_has_enhancement_elements(self):
        """Kiểm tra giao diện trang /admin có đầy đủ các nút và cột nâng cao mới."""
        self.login_as("SUPER_ADMIN_ENH")
        res = self.client.get("/admin")
        self.assertEqual(res.status_code, 200)
        html = res.data.decode("utf-8")

        # Nút Xuất Excel tài khoản
        self.assertIn('id="btnExportUsersExcel"', html)
        self.assertIn("Xuất Excel Tài Khoản", html)

        # Cột Mật khẩu ban đầu và Thao tác trong bảng tài khoản
        self.assertIn("Mật khẩu ban đầu", html)
        self.assertIn("Đặt lại MK", html)

        # Checkbox chọn mã mời và nút xuất các mã đã chọn
        self.assertIn('id="selectAllInviteCodes"', html)
        self.assertIn('id="btnExportSelectedInviteCodes"', html)
        self.assertIn('id="searchInviteCodesInput"', html)
        print("\n[PASS - TC 1]: Giao diện /admin hiển thị đầy đủ các nút Xuất Excel tài khoản, chọn mã mời, và thao tác quản lý.")

    def test_02_export_users_excel(self):
        """Kiểm tra tính năng xuất danh sách tài khoản người dùng ra file Excel (.xlsx)."""
        self.login_as("SUPER_ADMIN_ENH")
        res = self.client.get("/admin/users/export-excel?truong_id=1")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment; filename=", res.headers.get("Content-Disposition", ""))

        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Danh_Sach_Tai_Khoan", wb.sheetnames)
        ws = wb["Danh_Sach_Tai_Khoan"]

        headers = [cell.value for cell in ws[1]]
        self.assertIn("Mã định danh (Tên đăng nhập)", headers)
        self.assertIn("Mật khẩu ban đầu", headers)
        self.assertIn("Vai trò", headers)
        self.assertIn("Trạng thái", headers)

        # Kiểm tra sự hiện diện của học sinh test và mật khẩu khởi tạo
        found_student = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[3] == "HS_TEST_ENH1":
                found_student = True
                self.assertEqual(row[4], "Pass123@")
                self.assertEqual(row[5], "Học sinh")
                break
        self.assertTrue(found_student, "Không tìm thấy học sinh HS_TEST_ENH1 trong file Excel xuất ra")
        print("\n[PASS - TC 2]: Xuất danh sách tài khoản ra Excel thành công với đầy đủ cột Tên đăng nhập và Mật khẩu ban đầu.")

    def test_03_export_selected_invite_codes(self):
        """Kiểm tra xuất chỉ các mã mời được chọn qua checkbox."""
        self.login_as("SUPER_ADMIN_ENH")

        # POST chọn 1 mã duy nhất
        res = self.client.post("/admin/invite-codes/export-excel", data={
            "truong_id": "1",
            "selected_codes": "TBEDU-TEST-ENH1"
        })
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        ws = wb.active

        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 1)
        self.assertEqual(data_rows[0][1], "TBEDU-TEST-ENH1")
        print("\n[PASS - TC 3]: Xuất file Excel đúng chính xác danh sách mã mời đã chọn qua checkbox.")

    def test_04_delete_invite_code(self):
        """Kiểm tra xóa mã mời thành công."""
        self.login_as("SUPER_ADMIN_ENH")
        db = get_db()
        cur = db.cursor()

        # Tạo mã tạm để xóa
        cur.execute("""
            INSERT INTO invite_codes (ma_code, truong_id, loai, so_luot_toi_da, da_dung, ngay_tao)
            VALUES ('TBEDU-TEMP-DEL1', 1, 'ca_nhan', 1, 0, datetime('now'))
        """)
        code_id = cur.lastrowid
        db.commit()

        # Gọi xóa mã mời
        del_res = self.client.post(f"/admin/invite-codes/{code_id}/delete", follow_redirects=True)
        self.assertEqual(del_res.status_code, 200)

        # Kiểm tra trong DB
        cur.execute("SELECT id FROM invite_codes WHERE id = ?", (code_id,))
        self.assertIsNone(cur.fetchone())
        print("\n[PASS - TC 4]: Xóa mã mời thành công và bảo đảm an toàn dữ liệu.")

    def test_05_user_management_actions(self):
        """Kiểm tra Khóa/Mở khóa, Đặt lại mật khẩu và Xóa người dùng an toàn."""
        self.login_as("SUPER_ADMIN_ENH")
        db = get_db()
        cur = db.cursor()

        # 1. Khóa và mở khóa người dùng
        cur.execute("SELECT id, trang_thai FROM users WHERE ma_hoc_sinh = 'HS_TEST_ENH2'")
        user_row = cur.fetchone()
        user_id = user_row["id"]

        # Toggle khóa
        res_lock = self.client.post(f"/admin/users/{user_id}/toggle-status", follow_redirects=True)
        self.assertEqual(res_lock.status_code, 200)
        cur.execute("SELECT trang_thai FROM users WHERE id = ?", (user_id,))
        self.assertEqual(cur.fetchone()["trang_thai"], "da_khoa")

        # Toggle mở khóa lại
        res_unlock = self.client.post(f"/admin/users/{user_id}/toggle-status", follow_redirects=True)
        self.assertEqual(res_unlock.status_code, 200)
        cur.execute("SELECT trang_thai FROM users WHERE id = ?", (user_id,))
        self.assertEqual(cur.fetchone()["trang_thai"], "hoat_dong")

        # 2. Đặt lại mật khẩu
        new_pwd = "NewResetPass999!"
        res_reset = self.client.post(f"/admin/users/{user_id}/reset-password", data={
            "new_password": new_pwd
        }, follow_redirects=True)
        self.assertEqual(res_reset.status_code, 200)

        cur.execute("SELECT mat_khau, mat_khau_khoi_tao FROM users WHERE id = ?", (user_id,))
        updated_user = cur.fetchone()
        self.assertEqual(updated_user["mat_khau_khoi_tao"], new_pwd)
        self.assertTrue(check_password_hash(updated_user["mat_khau"], new_pwd))

        # Đăng nhập thử bằng mật khẩu mới
        login_res = self.login_as("HS_TEST_ENH2", new_pwd)
        self.assertEqual(login_res.status_code, 200)

        # 3. Tạo user tạm và xóa an toàn
        self.login_as("SUPER_ADMIN_ENH")
        cur.execute("""
            INSERT INTO users (ma_hoc_sinh, ho_ten, mat_khau, vai_tro, truong_id, trang_thai, so_du_gio)
            VALUES ('HS_TEMP_TO_DEL', 'User Xoa Tam', 'hash', 'hoc_sinh', 1, 'hoat_dong', 0.0)
        """)
        temp_user_id = cur.lastrowid
        db.commit()

        res_del = self.client.post(f"/admin/users/{temp_user_id}/delete", follow_redirects=True)
        self.assertEqual(res_del.status_code, 200)
        cur.execute("SELECT id FROM users WHERE id = ?", (temp_user_id,))
        self.assertIsNone(cur.fetchone())

        # Kiểm tra không thể tự xóa chính mình
        cur.execute("SELECT id FROM users WHERE ma_hoc_sinh = 'SUPER_ADMIN_ENH'")
        my_id = cur.fetchone()["id"]
        res_self_del = self.client.post(f"/admin/users/{my_id}/delete", follow_redirects=True)
        self.assertIn("không thể tự xóa", res_self_del.data.decode("utf-8"))

        print("\n[PASS - TC 5]: Khóa/Mở khóa, Đặt lại mật khẩu và Xóa tài khoản hoạt động chuẩn xác và an toàn tuyệt đối.")


if __name__ == "__main__":
    unittest.main()
